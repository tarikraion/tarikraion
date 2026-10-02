"""Merge research JSON into accounts.json, contacts CSV/XLSX, and the ICP page."""
import csv
import json
import re
from pathlib import Path

SCRATCH = Path("/tmp/claude-0/-home-user/022a5bf0-aa1e-54eb-849d-7d632705af27/scratchpad")
OUT = SCRATCH / "icp"
SOURCES = {
    "US": SCRATCH / "accounts_us.json",
    "UK": SCRATCH / "accounts_uk.json",
    "EU": SCRATCH / "accounts_eu.json",
    "ENT": SCRATCH / "accounts_enterprise.json",
}
EU_COUNTRIES = {
    "germany", "france", "netherlands", "sweden", "norway", "denmark", "finland",
    "switzerland", "spain", "italy", "poland", "estonia", "latvia", "lithuania",
    "austria", "belgium", "portugal", "czechia", "czech republic", "ireland",
    "luxembourg", "iceland", "greece", "romania", "bulgaria", "hungary", "slovenia",
    "croatia", "slovakia", "the netherlands",
}
C_LEVEL = re.compile(
    r"\b(ceo|cto|cfo|coo|cio|cdo|caio|cso|president|chief|founder|co-founder|cofounder|"
    r"director general|managing director|executive chair)\b",
    re.I,
)



SECTORS = {
    "Model labs": ["General Intuition", "Arcee", "Mirendil", "Odyssey", "Cognition", "Emulate", "Cosine", "Ineffable",
                   "AMI Labs", "H Company", "Cohere", "Pleias", "Mistral"],
    "Generative media & voice": ["Suno", "Higgsfield", "Fish Audio", "Runway", "Reactor", "Synthesia", "Stability",
                                 "ElevenLabs", "PolyAI", "Black Forest", "Gradium", "Wispr"],
    "Robotics & physical AI": ["Generalist", "Physical Intelligence", "Mind Robotics", "Wayve", "Humanoid", "NEURA",
                               "Gravis", "THEKER", "mimic"],
    "AI for science & bio": ["Chai", "Periodic", "Aureka", "Lila", "Axiom", "Harmonic", "Xaira", "CuspAI", "PhysicsX",
                             "Prima Mente", "Basecamp", "Bioptimus", "Cradle"],
    "Defence AI": ["Armadin", "Helsing", "Quantum Systems", "Harmattan", "STARK", "Alta Ares"],
    "AI apps & inference": ["Instinct", "Doubleword", "DeepL", "Lovable"],
    "Quant & trading": ["Jane Street", "Hudson River", "XTX", "Citadel", "Man Group"],
    "Pharma & banks": ["Bristol Myers", "Roche", "Eli Lilly", "HSBC", "Lloyds"],
    "Telco & sovereign cloud": ["Deutsche Telekom", "Telefónica", "Orange", "Telenor", "Proximus", "BT Group", "Nokia"],
    "Public sector & research": ["UK Department", "University of Edinburgh"],
    "Power & site owners": ["EDF", "ACS", "Schwarz"],
}
ICP_OF_SECTOR = {
    "Model labs": 1, "Generative media & voice": 1, "Robotics & physical AI": 1, "AI for science & bio": 1,
    "Defence AI": 1, "AI apps & inference": 2, "Quant & trading": 3, "Pharma & banks": 3,
    "Telco & sovereign cloud": 3, "Public sector & research": 3, "Power & site owners": 4,
}
ICP_OVERRIDE = {"Higgsfield": 2, "Wispr": 2, "ElevenLabs": 2, "PolyAI": 2, "Synthesia": 2, "Reactor": 2, "Cognition": 2}
ICP_NAMES = {1: "Scaling model builders", 2: "Inference-heavy products", 3: "Sovereign & regulated", 4: "Capital & power owners"}


def classify(company):
    for sector, names in SECTORS.items():
        for n in names:
            if company.lower().startswith(n.lower()):
                icp = ICP_OVERRIDE.get(n, ICP_OF_SECTOR[sector])
                return sector, icp
    raise KeyError(f"no sector for {company}")


def region_of(country: str) -> str:
    c = (country or "").strip().lower()
    if c in {"usa", "us", "united states", "united states of america"}:
        return "US"
    if c in {"uk", "united kingdom", "england", "scotland", "wales", "great britain"}:
        return "UK"
    if c in EU_COUNTRIES:
        return "Europe"
    return "Other"


def tier_of(score: int) -> str:
    return "A" if score >= 8 else "B" if score >= 6 else "C"


def load():
    seen = {}
    for src, path in SOURCES.items():
        if not path.exists():
            print(f"missing {path.name}")
            continue
        rows = json.loads(path.read_text())
        for r in rows:
            key = re.sub(r"[^a-z0-9]", "", r["company"].lower())
            r["source_set"] = src
            r["region"] = region_of(r.get("hq_country", ""))
            r["hq_city"] = re.sub(r"\s*\([^)]*\)", "", r.get("hq_city") or "").strip(" ,;")
            if r["hq_city"] in ("UK", "Unverified"):
                r["hq_city"] = ""
            r["fit_score"] = int(r.get("fit_score") or 0)
            r["tier"] = tier_of(r["fit_score"])
            r["sector"], r["icp"] = classify(r["company"])
            r["icp_name"] = ICP_NAMES[r["icp"]]
            r["contacts"] = [
                c for c in r.get("contacts", []) if C_LEVEL.search(c.get("title", ""))
            ]
            r["signals"] = sorted(
                r.get("signals", []), key=lambda s: s.get("date") or "", reverse=True
            )
            if key in seen:
                # keep the richer record
                if len(r["signals"]) + len(r["contacts"]) <= len(seen[key]["signals"]) + len(seen[key]["contacts"]):
                    continue
            seen[key] = r
    rows = sorted(seen.values(), key=lambda r: (-r["fit_score"], r["company"].lower()))
    return rows


def write_contacts(rows):
    header = [
        "first_name", "last_name", "full_name", "title", "company", "website",
        "linkedin_url", "hq_country", "region", "segment", "icp_segment", "tier", "fit_score",
        "top_signal", "signal_date", "signal_source", "contact_source", "confidence",
    ]
    out = []
    for r in rows:
        top = r["signals"][0] if r["signals"] else {}
        for c in r["contacts"]:
            parts = c["name"].split()
            out.append({
                "first_name": parts[0] if parts else "",
                "last_name": " ".join(parts[1:]),
                "full_name": c["name"],
                "title": c.get("title", ""),
                "company": r["company"],
                "website": r.get("website", ""),
                "linkedin_url": c.get("linkedin_url") or "",
                "hq_country": r.get("hq_country", ""),
                "region": r["region"],
                "segment": r["sector"],
                "icp_segment": r["icp_name"],
                "tier": r["tier"],
                "fit_score": r["fit_score"],
                "top_signal": top.get("signal", ""),
                "signal_date": top.get("date", ""),
                "signal_source": top.get("source_url", ""),
                "contact_source": c.get("source_url", ""),
                "confidence": c.get("confidence", ""),
            })
    with open(OUT / "raion_icp_c_level_contacts.csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=header)
        w.writeheader()
        w.writerows(out)
    return header, out


def write_xlsx(rows, header, contacts):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Accounts"
    acc_header = [
        "Company", "Website", "HQ city", "HQ country", "Region", "Sector", "ICP segment", "Tier",
        "Fit score", "What they do", "Top signal", "Signal date", "Signal source",
        "Est. GPU need", "Why Feb-Apr 2027", "C-level contacts",
    ]
    ws.append(acc_header)
    for r in rows:
        top = r["signals"][0] if r["signals"] else {}
        ws.append([
            r["company"], r.get("website", ""), r.get("hq_city", ""), r.get("hq_country", ""),
            r["region"], r["sector"], r["icp_name"], r["tier"], r["fit_score"],
            r.get("what_they_do", ""), top.get("signal", ""), top.get("date", ""),
            top.get("source_url", ""), r.get("est_gpu_need", ""), r.get("timing_rationale", ""),
            "; ".join(f'{c["name"]} ({c.get("title","")})' for c in r["contacts"]),
        ])
    ws2 = wb.create_sheet("C-level contacts")
    ws2.append(header)
    for c in contacts:
        ws2.append([c[h] for h in header])
    for sheet, widths in ((ws, [24, 26, 14, 14, 9, 22, 22, 6, 8, 44, 50, 11, 40, 34, 44, 50]),
                          (ws2, [12, 16, 22, 30, 24, 26, 40, 12, 9, 22, 22, 6, 8, 50, 11, 40, 40, 10])):
        for i, wdt in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(i)].width = wdt
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="000000")
            cell.alignment = Alignment(vertical="center")
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = sheet.dimensions
    wb.save(OUT / "raion_icp_target_accounts.xlsx")


def main():
    rows = load()
    (OUT / "accounts.json").write_text(json.dumps(rows, indent=1))
    header, contacts = write_contacts(rows)
    write_xlsx(rows, header, contacts)
    tpl = (OUT / "template.html").read_text()
    market = OUT / "market.html"
    page = tpl.replace("/*__ACCOUNTS__*/[]", json.dumps(rows).replace("</", "<\\/"))
    page = page.replace("<!--__MARKET_SECTION__-->", market.read_text() if market.exists() else "")
    (OUT / "raion-gpu-icp.html").write_text(page)
    by_region = {}
    for r in rows:
        by_region[r["region"]] = by_region.get(r["region"], 0) + 1
    print(f"accounts={len(rows)} contacts={len(contacts)} by_region={by_region}")
    print("tiers", {t: sum(1 for r in rows if r["tier"] == t) for t in "ABC"})


if __name__ == "__main__":
    main()
